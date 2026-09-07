import sys
from pptx import Presentation

def extract_text(filename):
    prs = Presentation(filename)
    text = []
    for i, slide in enumerate(prs.slides):
        text.append(f'--- Slide {i+1} ---')
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                text.append(shape.text)
    return '\n'.join(text)

if __name__ == '__main__':
    filename = 'Copie de Copie de Anatomy & Physiology Case Report by Slidesgo.pptx.pptx'
    try:
        print(extract_text(filename))
    except Exception as e:
        print(f"Error reading pptx: {e}")
